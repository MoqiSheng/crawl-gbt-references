#!/usr/bin/env python3
"""Append a prepared bibliography to Markdown, text, or DOCX without overwriting by default."""

from __future__ import annotations

import argparse
import re
from pathlib import Path


NUMBER_RE = re.compile(r"^\s*\[\d+\]\s*")
HEADING_RE = re.compile(r"^\s*#{1,6}\s*参考文献\s*$")


def read_references(path: Path) -> list[str]:
    references = []
    for line in path.read_text(encoding="utf-8-sig").splitlines():
        value = line.strip()
        if not value or value.startswith("#") or value.startswith("【"):
            continue
        value = NUMBER_RE.sub("", value)
        if value:
            references.append(value)
    if not references:
        raise SystemExit(f"No bibliography entries found in {path}")
    return references


def default_output(article: Path) -> Path:
    return article.with_name(f"{article.stem}_with_references{article.suffix}")


def ensure_output_available(output: Path, overwrite: bool) -> None:
    if output.exists() and not overwrite:
        raise SystemExit(f"Output already exists: {output}. Pass --overwrite to replace it.")
    output.parent.mkdir(parents=True, exist_ok=True)


def insert_text(article: Path, output: Path, references: list[str], replace_existing: bool) -> None:
    text = article.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    existing = next((i for i, line in enumerate(lines) if HEADING_RE.match(line)), None)
    if existing is not None:
        if not replace_existing:
            raise SystemExit("A 参考文献 heading already exists. Pass --replace-existing to replace that section.")
        lines = lines[:existing]
    body = "\n".join(lines).rstrip()
    block = "# 参考文献\n\n" + "\n\n".join(references) + "\n"
    output.write_text((body + "\n\n" if body else "") + block, encoding="utf-8")


def remove_from_paragraph(document, start_paragraph) -> None:
    body = document._body._element
    start = start_paragraph._element
    deleting = False
    for child in list(body):
        if child is start:
            deleting = True
        if deleting and child.tag.rsplit("}", 1)[-1] != "sectPr":
            body.remove(child)


def insert_docx(article: Path, output: Path, references: list[str], replace_existing: bool) -> None:
    try:
        from docx import Document
        from docx.shared import Cm
    except ImportError as error:
        raise SystemExit("python-docx is required. Run scripts/setup.ps1 first.") from error

    document = Document(article)
    existing = next((p for p in document.paragraphs if p.text.strip() == "参考文献"), None)
    if existing is not None:
        if not replace_existing:
            raise SystemExit("A 参考文献 heading already exists. Pass --replace-existing to replace that section.")
        remove_from_paragraph(document, existing)

    document.add_heading("参考文献", level=1)
    for reference in references:
        paragraph = document.add_paragraph(reference)
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description="Append GB/T references to a manuscript.")
    parser.add_argument("article", type=Path)
    parser.add_argument("bibliography", type=Path)
    parser.add_argument("--out", type=Path)
    parser.add_argument("--replace-existing", action="store_true")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    if not args.article.exists():
        raise SystemExit(f"Article not found: {args.article}")
    if not args.bibliography.exists():
        raise SystemExit(f"Bibliography not found: {args.bibliography}")
    if args.article.suffix.lower() == ".doc":
        raise SystemExit("Legacy .doc is not supported. Save it as .docx first.")

    output = args.out or default_output(args.article)
    ensure_output_available(output, args.overwrite)
    references = read_references(args.bibliography)
    suffix = args.article.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        insert_text(args.article, output, references, args.replace_existing)
    elif suffix == ".docx":
        insert_docx(args.article, output, references, args.replace_existing)
    else:
        raise SystemExit("Supported article types: .md, .markdown, .txt, .docx")

    print(f"Inserted {len(references)} references into {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
